"""Lote 3B · descargadores del Mac (fetch_nzd_b2, fetch_chf_snb, fetch_tona_mac): original frente a nuevo.

Sin red: las respuestas se construyen con el formato de cada fuente (XLSX B2 del RBNZ generado con openpyxl a
partir de los CSV reales del repo; CSV del portal SNB; RSS del SNB; JSON del BoJ). El original usa su
`requests` (curl_cffi no disponible → misma alternativa que en un Mac sin curl_cffi); el nuevo, el transporte
de g8common con las mismas respuestas. Casos: equivalencia, 503 con reintentos y copia local intacta, respuesta
vacía, serie inválida (las demás se escriben), fuente caída (CHF: las otras siguen), interrupción atómica,
Retry-After persistido, recuperación y registro state/fetch_<clave>.json → latido → aviso con el motivo."""
import csv
import importlib.util
import io
import json
import os
import shutil
import sys
import tempfile
import time
import unittest
from contextlib import redirect_stderr, redirect_stdout
from datetime import datetime, timedelta
from unittest import mock

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "scripts"))
sys.path.insert(0, os.path.dirname(__file__))

import fetcher_harness as H  # noqa: E402
from g8common import macfetch, series as S  # noqa: E402

NZ_IDS = {"INM.DB01.NZZV": "NZD_BILL_30D.csv", "INM.DB02.NZZV": "NZD_BILL_60D.csv", "INM.DB03.NZZV": "NZD_BILL_90D.csv",
          "INM.DG101.NZZCF": "NZD_BOND_1Y.csv", "INM.DG102.NZZCF": "NZD_BOND_2Y.csv", "INM.DG105.NZZCF": "NZD_BOND_5Y.csv",
          "INM.DG110.NZZCF": "NZD_BOND_10Y.csv", "INM.DN.NZK": "NZD_CASH_ON.csv"}
NZ_IIB = ["NZD_IIB_2035.csv", "NZD_IIB_2040.csv"]
CHF_T = {"1J": 1, "2J": 2, "3J": 3, "4J": 4, "5J": 5, "6J": 6, "7J": 7, "8J": 8, "9J": 9, "10J": 10, "20J": 20, "30J": 30}
_cnt = [0]


def dv_rows(fname, source=False):
    with open(os.path.join(ROOT, "data", fname), encoding="utf-8") as fh:
        rd = list(csv.DictReader(fh))
    return [(r["Date"], r["Value"]) + ((r.get("Source") or "",) if source else ()) for r in rd]


def nz_workbook(data, extra_date=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    cols = list(NZ_IDS.items()) + [("INM.DI%s.NZZC" % f[8:12], f) for f in NZ_IIB]
    ws.cell(1, 1, "Series")
    for j, (sid, f) in enumerate(cols, start=2):
        label = ("Inflation indexed bonds" if "IIB" in f else "Wholesale rates")
        ws.cell(1, j, label)
        ws.cell(2, j, ("20 September %s" % f[8:12]) if "IIB" in f else f[:-4])
        ws.cell(5, j, sid)
    look = {f: dict(data[f]) for _, f in cols}
    dates = sorted({d for f in look for d in look[f]})
    for i, d in enumerate(dates, start=6):
        ws.cell(i, 1, datetime.strptime(d, "%Y-%m-%d"))
        for j, (_, f) in enumerate(cols, start=2):
            if d in look[f]:
                ws.cell(i, j, float(look[f][d]))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def load(path, name):
    _cnt[0] += 1
    spec = importlib.util.spec_from_file_location("%s_%d" % (name, _cnt[0]), path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class MacBase(object):
    key = script = orig_path = new_path = None
    files = ()

    def setUp(self):
        self.roots = []
        self.orig = {f: open(os.path.join(ROOT, "data", f), "rb").read() for f in self.files}

    def tearDown(self):
        macfetch.TEST_TRANSPORT = macfetch.TEST_CLOCK = None
        for r in self.roots:
            shutil.rmtree(r, ignore_errors=True)

    def root(self):
        r = tempfile.mkdtemp()
        os.makedirs(os.path.join(r, "data"))
        for f, b in self.orig.items():
            with open(os.path.join(r, "data", f), "wb") as fh:
                fh.write(b)
        self.roots.append(r)
        return r

    def run_orig(self, serve, root, argv=()):
        fake = H._FakeReqMod(serve)
        with mock.patch.dict(sys.modules, {"requests": fake, "curl_cffi": None}):   # también `import requests` local
            mod = load(self.orig_path, self.script + "_orig")
            return self._orig_inner(mod, fake, root, argv)

    def _orig_inner(self, mod, fake, root, argv):
        mod.requests = fake
        if hasattr(mod, "crequests"):
            mod.crequests, mod._IMPERSONATE = None, False
        if hasattr(mod, "time"):
            mod.time = H._NoSleep(mod.time)
        self.adjust(mod, root)
        return self._main(mod, root, argv), mod

    def run_new(self, serve, root, argv=(), clock=None):
        calls = []

        def transport(method, url, headers, body, ct, rt, tt):
            calls.append(url)
            st, data = serve(url)
            hdrs = {}
            if isinstance(data, tuple):
                data, hdrs = data
            return st, hdrs, data
        clock = clock or H.clock_at_now()
        macfetch.TEST_TRANSPORT, macfetch.TEST_CLOCK = transport, clock
        try:
            mod = load(self.new_path, self.script + "_new")
            self.adjust(mod, root)
            rc = self._main(mod, root, argv)
        finally:
            macfetch.TEST_TRANSPORT = macfetch.TEST_CLOCK = None
        return rc, calls, clock

    def adjust(self, mod, root):
        pass

    def _main(self, mod, root, argv):
        old, cwd = sys.argv, os.getcwd()
        sys.argv = [self.script + ".py"] + list(argv)
        os.chdir(root)
        try:
            with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                rc = mod.main()
        except SystemExit as e:
            rc = e.code
        finally:
            sys.argv = old
            os.chdir(cwd)
        return rc

    def read(self, root):
        return {f: H.read(root, f) for f in self.files}

    def state(self, root):
        return json.load(open(os.path.join(root, "state", "fetch_%s.json" % self.key)))

    # ── comunes ─────────────────────────────────────────────────────────────
    def test_E1_equivalence(self):
        serve = self.serve(new_day=True)
        ro, rn = self.root(), self.root()
        rco, _ = self.run_orig(serve, ro)
        rcn, calls, _ = self.run_new(serve, rn)
        self.assertEqual((rco, rcn), (0, 0))
        o, n = self.read(ro), self.read(rn)
        for f in self.files:
            self.assertEqual(n[f], o[f], f)
        self.assertTrue([f for f in self.files if n[f] != self.orig[f]])
        st = self.state(rn)
        self.assertEqual((st["rc"], st["errors"]), (0, []))

    def test_F1_503_retried_files_intact_recorded(self):
        rn = self.root()
        rc, calls, clock = self.run_new(lambda u: (503, b"busy"), rn)
        self.assertEqual(rc, 1)
        self.assertEqual(self.read(rn), self.orig)
        self.assertEqual(clock.slept[:3], [10, 40, 90])
        st = self.state(rn)
        self.assertEqual(st["rc"], 1)
        self.assertTrue(st["errors"])
        self.assertIn("FAIL_TRANSIENT", json.dumps(st["requests"]))

    def test_F2_waf_403_not_retried(self):
        rn = self.root()
        rc, calls, clock = self.run_new(lambda u: (403, b"blocked"), rn)
        self.assertEqual(rc, 1)
        self.assertFalse(getattr(clock, "slept", []))
        self.assertEqual(self.read(rn), self.orig)

    def test_F3_empty_response(self):
        rn = self.root()
        rc, *_ = self.run_new(lambda u: (200, b""), rn)
        self.assertNotEqual(rc, 0)                     # 1 = descarga, 2 = formato (código del original)
        self.assertEqual(self.read(rn), self.orig)

    def test_F6_interruption_atomic(self):
        rn = self.root()
        real = os.replace
        seen = []

        def flaky(src, dst):
            if dst.endswith(".csv"):
                seen.append(dst)
                if len(seen) == min(2, len(self.files)):
                    raise KeyboardInterrupt("corte simulado")
            return real(src, dst)
        with mock.patch("g8common.series.os.replace", flaky):
            with self.assertRaises(KeyboardInterrupt):
                self.run_new(self.serve(new_day=True), rn)
        got = self.read(rn)
        for f in self.files:
            S.parse(got[f])                                                     # ninguna a medias
        self.assertEqual(got[os.path.basename(seen[-1])], self.orig[os.path.basename(seen[-1])])

    def test_F7_retry_after_persisted_and_recovery(self):
        rn = self.root()
        clock = H.clock_at_now()
        rc, calls, _ = self.run_new(lambda u: (503, (b"", {"retry-after": "3600"})), rn, clock=clock)
        self.assertEqual(rc, 1)
        nb = json.load(open(os.path.join(rn, "data", "_ingest", "not_before", "mac.json")))
        self.assertTrue(nb["not_before"])
        clock.t += 600
        rc2, calls2, _ = self.run_new(self.serve(new_day=True), rn, clock=clock)
        self.assertEqual(calls2, [])                                            # no se consulta antes de tiempo
        clock.t += 3600
        rc3, calls3, _ = self.run_new(self.serve(new_day=True), rn, clock=clock)
        self.assertEqual(rc3, 0)                                                # recuperación
        self.assertNotEqual(self.read(rn), self.orig)
        self.assertEqual(self.state(rn)["errors"], [])


class Test_fetch_nzd_b2(MacBase, unittest.TestCase):
    key, script = "nzd", "fetch_nzd_b2"
    orig_path = os.path.join(ROOT, "tests", "fixtures", "original", "fetch_nzd_b2.py")
    new_path = os.path.join(ROOT, "scripts", "fetch_nzd_b2.py")
    files = tuple(NZ_IDS.values()) + tuple(NZ_IIB)
    _cache = {}

    def serve(self, new_day=False, drop=None):
        k = (new_day, drop)
        if k not in self._cache:
            data = {f: dv_rows(f) for f in self.files}
            if new_day:
                for f in self.files:
                    data[f] = data[f] + [("2026-09-24", data[f][-1][1])]
            if drop:
                data[drop] = []
            cur = nz_workbook(data)
            hist = nz_workbook({f: [] for f in self.files})
            self._cache[k] = (cur, hist)
        cur, hist = self._cache[k]
        return lambda u: (200, hist if "1985-2017" in u else cur)

    def test_F4_one_invalid_series_others_written(self):
        rn = self.root()
        rc, *_ = self.run_new(self.serve(new_day=True, drop="NZD_BOND_5Y.csv"), rn)
        self.assertEqual(rc, 1)
        got = self.read(rn)
        self.assertEqual(got["NZD_BOND_5Y.csv"], self.orig["NZD_BOND_5Y.csv"])     # vacía: no se escribe
        self.assertIn(b"2026-09-24", got["NZD_BOND_10Y.csv"])                      # las demás, sí
        st = self.state(rn)
        self.assertEqual(st["files"]["NZD_BOND_5Y.csv"]["status"], "REJECTED")
        ro = self.root()
        self.run_orig(self.serve(new_day=True, drop="NZD_BOND_5Y.csv"), ro)
        self.assertEqual(H.read(ro, "NZD_BOND_5Y.csv"), b"Date,Value\n")         # el original la vaciaba

    def test_F5_history_splice_failure_keeps_local_history(self):
        data = {f: dv_rows(f) for f in self.files}
        cut = {f: [r for r in rows if r[0] >= "2020-01-01"] + [("2026-09-24", rows[-1][1])] for f, rows in data.items()}
        cur = nz_workbook(cut)
        rn = self.root()
        rc, *_ = self.run_new(lambda u: (503, b"") if "1985-2017" in u else (200, cur), rn, argv=())
        got = self.read(rn)
        for f in self.files:
            old = S.parse(self.orig[f])
            new = S.parse(got[f])
            self.assertTrue(set(old.rows) <= set(new.rows), f)                  # ninguna fecha anterior se pierde
            self.assertEqual(new.max_date, "20260924", f)


class Test_fetch_chf_snb(MacBase, unittest.TestCase):
    key, script = "chf", "fetch_chf_snb"
    orig_path = os.path.join(ROOT, "tests", "fixtures", "original", "fetch_chf_snb.py")
    new_path = os.path.join(ROOT, "scripts", "fetch_chf_snb.py")
    files = tuple("CHF_SPOT_%dY.csv" % y for y in CHF_T.values()) + ("CHF_NOM_10Y.csv", "CHF_SARON.csv")

    def serve(self, new_day=False, fail=None):
        curve = ['"CubeId";"rendeiduebd"', '"PublishingDate";"2026-09-01 14:30"', "", '"Date";"D0";"D1";"Value"']
        for t, y in CHF_T.items():
            rows = dv_rows("CHF_SPOT_%dY.csv" % y)
            for d, v in rows[-60:]:
                curve.append('"%s";"CHF";"%s";"%s"' % (d, t, v))
        sar = ['"CubeId";"zirepo"', '"PublishingDate";"2026-09-23 09:00"', "", '"Date";"D0";"Value"']
        for d, v, src in dv_rows("CHF_SARON.csv", source=True)[-40:]:
            if src == "zirepo":
                sar.append('"%s";"H0";"%s"' % (d, v))
        rss = "<rss>" + "".join("<item><title>CH: %s %s %s SNB</title></item>" % (v, c, d) for c, d, v in (
            ("R10", "2026-09-23", "0.575"), ("SARH", "2026-09-22", "-0.04"), ("SNBLZ", "2026-09-22", "0.00"))
            + ((("R10", "2026-09-24", "0.58"), ("SARH", "2026-09-23", "-0.05")) if new_day else ())) + "</rss>"
        bodies = {"rendeiduebd": "\n".join(curve).encode(), "zirepo": "\n".join(sar).encode(), "rss": rss.encode()}

        def serve(u):
            k = "rendeiduebd" if "rendeiduebd" in u else "zirepo" if "zirepo" in u else "rss"
            if fail == k:
                return 503, b"busy"
            return 200, bodies[k]
        return serve

    def test_F4_one_source_down_others_written(self):
        ro, rn = self.root(), self.root()
        rco, _ = self.run_orig(self.serve(new_day=True, fail="zirepo"), ro)
        rcn, *_ = self.run_new(self.serve(new_day=True, fail="zirepo"), rn)
        self.assertEqual((rco, rcn), (1, 1))
        self.assertIn(b"2026-09-24,0.58,rss", H.read(rn, "CHF_NOM_10Y.csv"))       # 10Y del RSS: sí
        self.assertNotIn(b"2026-09-24", H.read(ro, "CHF_NOM_10Y.csv"))             # el original abortaba todo
        self.assertIn(b"2026-09-23,-0.05,rss", H.read(rn, "CHF_SARON.csv"))        # SARON del RSS: sí
        self.assertIn("saron", " ".join(self.state(rn)["errors"]))


class Test_fetch_tona_mac(MacBase, unittest.TestCase):
    key, script = "tona", "fetch_tona_mac"
    orig_path = os.path.join(ROOT, "tests", "fixtures", "original", "fetch_tona_mac.py")
    new_path = os.path.join(ROOT, "mac", "fetch_tona_mac.py")
    files = ("TONA.csv",)

    def adjust(self, mod, root):
        mod.OUT = os.path.join(root, "data", "TONA.csv")
        mod.datetime = H.fixed_dt(H.NOW)

    def serve(self, new_day=False):
        rows = [r for r in H.repo_rows("TONA.csv") if r[0] >= "20210901"]
        if new_day:
            rows = rows + [("20260924", rows[-1][1])]
        b = H.boj_json(rows)
        return lambda u: (200, b)

    def test_F4_boj_status_error_in_json_is_transient(self):
        rn = self.root()
        body = json.dumps({"STATUS": 503, "MESSAGE": "DB error"}).encode()
        rc, calls, clock = self.run_new(lambda u: (200, body), rn)
        self.assertEqual(rc, 1)
        self.assertEqual(len(calls), 4)                                            # 503 dentro del JSON: se reintenta
        self.assertEqual(self.read(rn), self.orig)


class FakeLibs(object):
    """Dobles de curl_cffi y requests para ejercitar el transporte REAL (macfetch.curl_transport) sin red.
    script: {etiqueta: [respuesta, …]} con etiqueta = perfil («safari17_0», «chrome124», «chrome») o «requests»;
    respuesta = (status, headers) | Exception | ("SLOW", segundos) (consume ese tiempo y agota el timeout).
    Cada llamada avanza el reloj simulado `cost` s y registra el timeout concedido."""

    def __init__(self, script, clock, cost=0.1, body=b"x"):
        import types
        self.script, self.clock, self.cost, self.body = script, clock, cost, body
        self.calls = []

        class R(object):
            def __init__(s2, st, hd, body):
                s2.status_code, s2.headers, s2.content = st, hd, body

        def serve(label, url, timeout, allow_redirects):
            self.calls.append({"label": label, "url": url, "timeout": timeout, "allow_redirects": allow_redirects,
                               "t": self.clock.t})
            seq = self.script.get(label) or [(403, {})]
            item = seq.pop(0) if len(seq) > 1 else seq[0]
            if isinstance(item, tuple) and item[0] == "SLOW":
                self.clock.t += min(item[1], timeout)
                raise OSError("Operation timed out")
            self.clock.t += self.cost
            if isinstance(item, Exception):
                raise item
            st, hd = item[0], item[1]
            body = item[2] if len(item) > 2 else self.body
            return R(st, hd, body)

        def creq(method, url, headers=None, data=None, timeout=None, impersonate=None, allow_redirects=True):
            return serve(impersonate, url, timeout, allow_redirects)

        def req(method, url, headers=None, data=None, timeout=None, allow_redirects=True):
            return serve("requests", url, timeout, allow_redirects)
        cc = types.ModuleType("curl_cffi")
        cc.requests = types.SimpleNamespace(request=creq)
        rq = types.ModuleType("requests")
        rq.request = req

        class Session(object):                          # el transporte usa Session + adaptador (B3R1-1)
            def request(s2, method, url, headers=None, data=None, timeout=None, allow_redirects=True):
                return req(method, url, headers=headers, data=data, timeout=timeout, allow_redirects=allow_redirects)

            def mount(s2, prefix, adapter):
                pass

            def close(s2):
                pass

        class HTTPAdapter(object):
            def __init__(s2, **kw):
                pass
        rq.Session = Session
        rq.adapters = types.SimpleNamespace(HTTPAdapter=HTTPAdapter)
        self.modules = {"curl_cffi": cc, "requests": rq}

    def labels(self):
        return [c["label"] for c in self.calls]


class CurlTransport(unittest.TestCase):
    """B3-2 / B3-3 — transporte real del Mac con las bibliotecas HTTP sustituidas por dobles."""

    def setUp(self):
        self.clock = H.clock_at_now()
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp)

    def fetch(self, libs, budget=900):
        with mock.patch.dict(sys.modules, libs.modules):
            F = macfetch.Fetch("t", self.tmp, budget_s=budget, now=self.clock, sleep=self.clock.sleep, log=lambda *a: None)
        return F

    def get(self, F, libs, url="https://www.stat-search.boj.or.jp/api/x"):
        with mock.patch.dict(sys.modules, libs.modules):
            return F.get(url, timeout=45)

    # selección legítima de transporte
    def test_first_200_wins_in_original_order_without_library_redirects(self):
        libs = FakeLibs({"safari17_0": [(200, {})]}, self.clock)
        self.assertEqual(self.get(self.fetch(libs), libs), b"x")
        self.assertEqual(libs.labels(), ["safari17_0"])
        self.assertFalse(libs.calls[0]["allow_redirects"])

    def test_waf_403_tries_next_profile(self):
        libs = FakeLibs({"safari17_0": [(403, {})], "chrome124": [(200, {})]}, self.clock)
        self.assertEqual(self.get(self.fetch(libs), libs), b"x")
        self.assertEqual(libs.labels(), ["safari17_0", "chrome124"])

    def test_all_profiles_blocked_is_fail_access_without_retries(self):
        libs = FakeLibs({}, self.clock)
        F = self.fetch(libs)
        with self.assertRaises(macfetch.FetchError) as cm:
            self.get(F, libs)
        self.assertEqual(cm.exception.cls, "FAIL_ACCESS")
        self.assertEqual(libs.labels(), ["safari17_0", "chrome124", "chrome", "requests"])
        self.assertFalse(getattr(self.clock, "slept", []))

    def test_network_error_tries_next_profile(self):
        libs = FakeLibs({"safari17_0": [OSError("TLS handshake")], "chrome124": [(200, {})]}, self.clock)
        self.assertEqual(self.get(self.fetch(libs), libs), b"x")

    # B3-2: Retry-After nunca queda oculto tras otro perfil
    def test_b3_2_retry_after_on_first_profile_single_request_persisted(self):
        self._retry_after_first_profile(429)

    def test_b3_2_503_with_retry_after_on_first_profile_single_request_persisted(self):
        self._retry_after_first_profile(503)

    def _retry_after_first_profile(self, status):
        libs = FakeLibs({"safari17_0": [(status, {"Retry-After": "7200"}), (200, {})], "chrome124": [(200, {})]},
                        self.clock)
        F = self.fetch(libs)
        with self.assertRaises(macfetch.FetchError) as cm:
            self.get(F, libs)
        self.assertEqual(cm.exception.cls, "DEFERRED")
        self.assertEqual(libs.labels(), ["safari17_0"])                            # una sola petición
        nb = json.load(open(os.path.join(self.tmp, "data", "_ingest", "not_before", "mac.json")))
        self.assertTrue(nb["not_before"])
        with self.assertRaises(macfetch.FetchError):                              # misma ejecución: bloqueado
            self.get(F, libs)
        self.clock.t += 3600
        F2 = self.fetch(libs)                                                      # ejecución siguiente: bloqueado
        with self.assertRaises(macfetch.FetchError) as cm2:
            self.get(F2, libs)
        self.assertEqual(cm2.exception.cls, "DEFERRED")
        self.assertEqual(libs.labels(), ["safari17_0"])
        self.clock.t += 3601                                                       # vencido: vuelve a consultar
        self.assertEqual(self.get(self.fetch(libs), libs), b"x")

    def test_b3_2_503_is_retried_by_g8http_not_masked_by_profiles(self):
        libs = FakeLibs({"safari17_0": [(503, {}), (503, {}), (200, {})], "chrome124": [(200, {})]}, self.clock)
        self.assertEqual(self.get(self.fetch(libs), libs), b"x")
        self.assertEqual(libs.labels(), ["safari17_0"] * 3)
        self.assertEqual(self.clock.slept, [10, 40])

    def test_b3_2_redirect_followed_by_g8http_within_budget(self):
        libs = FakeLibs({"safari17_0": [(302, {"Location": "https://www.stat-search.boj.or.jp/api/y"}), (200, {})]},
                        self.clock)
        self.assertEqual(self.get(self.fetch(libs), libs), b"x")
        self.assertEqual([c["url"][-1] for c in libs.calls], ["x", "y"])

    # B3-3: un solo plazo para toda la cadena
    def test_b3_3_chain_respects_total_budget(self):
        libs = FakeLibs({k: [("SLOW", 100)] for k in ("safari17_0", "chrome124", "chrome", "requests")}, self.clock)
        F = self.fetch(libs, budget=10)
        t0 = self.clock.t
        with self.assertRaises(macfetch.FetchError):
            self.get(F, libs)
        spent = self.clock.t - t0
        self.assertLessEqual(spent, 10 + 0.5)                                      # tolerancia explícita 0,5 s
        self.assertLessEqual(sum(c["timeout"] for c in libs.calls), 10 + 1e-6)     # concesiones = remanentes
        self.assertEqual(len(libs.calls), 1)                                       # nada tras agotar el plazo

    def test_b3_3_fast_alternative_after_slow_one_completes_in_budget(self):
        libs = FakeLibs({"safari17_0": [("SLOW", 6)], "chrome124": [(200, {})]}, self.clock)
        F = self.fetch(libs, budget=10)
        t0 = self.clock.t
        self.assertEqual(self.get(F, libs), b"x")
        self.assertLessEqual(self.clock.t - t0, 10)
        self.assertLessEqual(libs.calls[1]["timeout"], 10 - 6 + 1e-6)             # solo el remanente

    def test_b3_3_requests_fallback_gets_only_remaining(self):
        libs = FakeLibs({"safari17_0": [("SLOW", 3)], "chrome124": [("SLOW", 3)], "chrome": [("SLOW", 3)],
                         "requests": [(200, {})]}, self.clock)
        F = self.fetch(libs, budget=10)
        self.assertEqual(self.get(F, libs), b"x")
        self.assertLessEqual(libs.calls[-1]["timeout"], 1 + 1e-6)

    # integración de extremo a extremo con un descargador real del Mac
    def test_end_to_end_tona_retry_after_then_blocked_next_run(self):
        root = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, root)
        os.makedirs(os.path.join(root, "data"))
        orig = open(os.path.join(ROOT, "data", "TONA.csv"), "rb").read()
        open(os.path.join(root, "data", "TONA.csv"), "wb").write(orig)
        rows = [r for r in H.repo_rows("TONA.csv") if r[0] >= "20210901"] + [("20260924", "0.9770")]
        ok = (200, {}, H.boj_json(rows))
        libs = FakeLibs({"safari17_0": [(429, {"Retry-After": "7200"}), ok], "chrome124": [ok]}, self.clock)

        def run():
            macfetch.TEST_CLOCK = self.clock
            try:
                with mock.patch.dict(sys.modules, libs.modules):
                    mod = load(os.path.join(ROOT, "mac", "fetch_tona_mac.py"), "tona_e2e")
                    mod.OUT = os.path.join(root, "data", "TONA.csv")
                    mod.datetime = H.fixed_dt(H.NOW)
                    with redirect_stdout(io.StringIO()):
                        return mod.main()
            finally:
                macfetch.TEST_CLOCK = None
        self.assertEqual(run(), 1)
        self.assertEqual(libs.labels(), ["safari17_0"])
        self.assertEqual(H.read(root, "TONA.csv"), orig)
        st = json.load(open(os.path.join(root, "state", "fetch_tona.json")))
        self.assertIn("DEFERRED", " ".join(st["errors"]))
        self.clock.t += 600
        self.assertEqual(run(), 1)                                                 # siguiente ejecución: sin peticiones
        self.assertEqual(libs.labels(), ["safari17_0"])
        self.clock.t += 7200
        self.assertEqual(run(), 0)                                                 # vencido: descarga y escribe
        self.assertIn(b"20260924", H.read(root, "TONA.csv"))
        st = json.load(open(os.path.join(root, "state", "fetch_tona.json")))
        self.assertEqual(st["files"]["TONA.csv"]["max_date"], "20260924")         # B3-4 en el registro real


class _Slow(object):
    """Servidor HTTP local (127.0.0.1) para B3R1-1. mode:
       drip       — cabeceras al momento, cuerpo de 10 B con Content-Length, 1 B cada 0,30 s;
       drip_nolen — igual pero sin Content-Length (fin por cierre): un corte parecería una respuesta completa;
       headers    — la línea de estado y las cabeceras llegan 1 B cada 0,10 s;
       fast       — respuesta completa inmediata."""

    def __init__(self, mode, tls=None):
        import http.server
        import threading
        outer = self
        self.hits, self.aborted = 0, threading.Event()

        class Handler(http.server.BaseHTTPRequestHandler):
            protocol_version = "HTTP/1.0"

            def log_message(self, *a):
                pass

            def do_GET(self):
                outer.hits += 1
                try:
                    if mode == "headers":
                        for b in b"HTTP/1.0 200 OK\r\nContent-Length: 1\r\n\r\nx":
                            time.sleep(0.10)
                            self.wfile.write(bytes([b]))
                            self.wfile.flush()
                        return
                    self.send_response(200)
                    if mode != "drip_nolen":
                        self.send_header("Content-Length", "10")
                    self.end_headers()
                    if mode == "fast":
                        self.wfile.write(b"x" * 10)
                        return
                    for _ in range(10):
                        time.sleep(0.30)
                        self.wfile.write(b"x")
                        self.wfile.flush()
                except OSError:                         # EPIPE/ECONNRESET o, con TLS, SSLEOFError
                    outer.aborted.set()                 # el cliente cortó la conexión: recursos liberados

        self.server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        if tls:                                         # (cert, key): TLS real, verificado por el cliente
            import ssl
            ctx = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
            ctx.load_cert_chain(*tls)
            self.server.socket = ctx.wrap_socket(self.server.socket, server_side=True)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = "%s://127.0.0.1:%d/x" % ("https" if tls else "http", self.server.server_port)

    def close(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()


class _Proxy(object):
    """Proxy HTTP local (127.0.0.1) para B3R2-1. mode:
       slow_connect — responde al CONNECT «HTTP/1.0 502 Bad Gateway» de 1 B cada 0,10 s (no toca el destino);
       tunnel       — CONNECT normal: abre el túnel al destino (solo 127.0.0.1) y reenvía en ambos sentidos."""

    def __init__(self, mode):
        import select
        import socket
        import socketserver
        import threading
        outer = self
        self.hits, self.aborted = 0, threading.Event()

        class H(socketserver.StreamRequestHandler):
            def handle(self):
                line = self.rfile.readline()
                while self.rfile.readline() not in (b"\r\n", b"\n", b""):
                    pass
                outer.hits += 1
                try:
                    if mode == "slow_connect":
                        for b in b"HTTP/1.0 502 Bad Gateway\r\n\r\n":
                            time.sleep(0.10)
                            self.wfile.write(bytes([b]))
                            self.wfile.flush()
                        return
                    host, port = line.split()[1].decode().rsplit(":", 1)
                    assert host == "127.0.0.1", host                # nunca fuera de la máquina
                    up = socket.create_connection((host, int(port)))
                    self.wfile.write(b"HTTP/1.0 200 Connection established\r\n\r\n")
                    self.wfile.flush()
                    a = self.connection
                    while True:
                        r, _, _ = select.select([a, up], [], [], 5)
                        if not r:
                            break
                        for src in r:
                            data = src.recv(65536)
                            if not data:
                                up.close()
                                return
                            (up if src is a else a).sendall(data)
                except OSError:
                    outer.aborted.set()

        class Srv(socketserver.ThreadingTCPServer):
            daemon_threads = True
            allow_reuse_address = True
        self.server = Srv(("127.0.0.1", 0), H)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = "http://127.0.0.1:%d" % self.server.server_address[1]

    def close(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()


class _SlowTLSHandshake(object):
    """Servidor TCP local que, en lugar del ServerHello, envía la cabecera de un registro TLS de 16 KiB y luego su
    contenido de 1 B cada 0,10 s: el cliente queda dentro del saludo TLS (connect()) sin agotar ninguna espera."""

    def __init__(self):
        import socketserver
        import threading
        outer = self
        self.hits, self.aborted = 0, threading.Event()

        class H(socketserver.BaseRequestHandler):
            def handle(self):
                outer.hits += 1
                try:
                    self.request.recv(4096)                          # ClientHello
                    self.request.sendall(b"\x16\x03\x03\x40\x00")
                    for _ in range(100):
                        time.sleep(0.10)
                        self.request.sendall(b"\x00")
                except OSError:
                    outer.aborted.set()

        class Srv(socketserver.ThreadingTCPServer):
            daemon_threads = True
            allow_reuse_address = True
        self.server = Srv(("127.0.0.1", 0), H)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = "https://127.0.0.1:%d/x" % self.server.server_address[1]

    def close(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()


try:
    import requests as _real_requests  # noqa: F401
    HAS_REQUESTS = True
except ImportError:
    HAS_REQUESTS = False
NEEDS_REQUESTS = unittest.skipUnless(HAS_REQUESTS, "requests no instalado: la alternativa requests no existe aquí")


class RequestsDeadlineReal(unittest.TestCase):
    """B3R1-1 — plazo total EFECTIVO con requests REAL (y la biblioteca estándar) frente a un servidor lento local.
    Reloj real; solo 127.0.0.1; sin proveedores. El cierre del servidor se hace después de medir la llamada."""
    BUDGET, TOL = 1.2, 0.5

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp)
        env = mock.patch.dict(os.environ, {"NO_PROXY": "127.0.0.1,localhost", "no_proxy": "127.0.0.1,localhost"})
        env.start()
        self.addCleanup(env.stop)

    def call(self, mode, stdlib=False, tls=None):
        import time as _t
        srv = _Slow(mode, tls)
        mods = {"curl_cffi": None}
        if stdlib:
            mods["requests"] = None                     # alternativa final: g8http.default_transport
        try:
            with mock.patch.dict(sys.modules, mods):
                F = macfetch.Fetch("t", self.tmp, budget_s=self.BUDGET, log=lambda *a: None)
                t0 = _t.monotonic()
                try:
                    body, err = F.get(srv.url, timeout=120), None
                except macfetch.FetchError as e:
                    body, err = None, e
                elapsed = _t.monotonic() - t0
            aborted = srv.aborted.wait(3) if err is not None else None
            return body, err, elapsed, F, srv.hits, aborted
        finally:
            srv.close()

    def assert_cut(self, mode, stdlib=False, tls=None):
        body, err, elapsed, F, hits, aborted = self.call(mode, stdlib, tls)
        self.assertLessEqual(elapsed, self.BUDGET + self.TOL, "%s: %.3f s" % (mode, elapsed))
        self.assertIsNone(body)                                         # nunca OK fuera de plazo
        self.assertIsNotNone(err)
        self.assertNotEqual(F.requests_log[-1]["cls"], "OK")
        self.assertEqual(hits, 1)                                       # ningún transporte ni reintento tras agotarlo
        self.assertTrue(aborted, "la conexión no se cortó")              # recursos liberados

    @NEEDS_REQUESTS
    def test_requests_real_drip_body_cut_at_deadline(self):
        self.assert_cut("drip")

    @NEEDS_REQUESTS
    def test_requests_real_drip_without_length_not_taken_as_complete(self):
        self.assert_cut("drip_nolen")

    @NEEDS_REQUESTS
    def test_requests_real_slow_headers_cut_at_deadline(self):
        self.assert_cut("headers")

    def test_stdlib_fallback_drip_body_cut_at_deadline(self):
        self.assert_cut("drip", stdlib=True)

    def test_stdlib_fallback_slow_headers_cut_at_deadline(self):
        self.assert_cut("headers", stdlib=True)

    def test_fast_response_still_ok(self):
        for stdlib in (False, True):
            body, err, elapsed, F, hits, _ = self.call("fast", stdlib)
            self.assertIsNone(err, stdlib)
            self.assertEqual(body, b"x" * 10)
            self.assertEqual(F.requests_log[-1]["cls"], "OK")
            self.assertLess(elapsed, self.BUDGET)

    # HTTPS (los proveedores del Mac lo son): certificado autofirmado para 127.0.0.1, VERIFICADO por el cliente
    # (REQUESTS_CA_BUNDLE para requests, SSL_CERT_FILE para la biblioteca estándar); nunca se desactiva TLS.
    def tls(self):
        import subprocess
        d = os.path.join(self.tmp, "tls")
        os.makedirs(d)
        cert, key = os.path.join(d, "c.pem"), os.path.join(d, "k.pem")
        try:
            subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-days", "1", "-subj",
                            "/CN=127.0.0.1", "-addext", "subjectAltName=IP:127.0.0.1", "-keyout", key, "-out", cert],
                           check=True, capture_output=True)
        except (OSError, subprocess.CalledProcessError) as e:
            self.skipTest("openssl no disponible: %s" % e)
        env = mock.patch.dict(os.environ, {"REQUESTS_CA_BUNDLE": cert, "SSL_CERT_FILE": cert})
        env.start()
        self.addCleanup(env.stop)
        return cert, key

    @NEEDS_REQUESTS
    def test_https_requests_real_drip_cut_at_deadline(self):
        self.assert_cut("drip", tls=self.tls())

    def test_https_stdlib_fallback_drip_cut_at_deadline(self):
        self.assert_cut("drip", stdlib=True, tls=self.tls())

    def test_https_fast_response_still_ok(self):
        t = self.tls()
        for stdlib in (False, True):
            body, err, *_ = self.call("fast", stdlib, t)
            self.assertIsNone(err, stdlib)
            self.assertEqual(body, b"x" * 10)

    @NEEDS_REQUESTS
    # ── B3R2-1: fases dentro de connect() — túnel CONNECT de un proxy HTTP y saludo TLS ──
    def fetch_via(self, target_url, proxy=None, stdlib=False):
        import time as _t
        env = {}                                        # sin proxy: el NO_PROXY de setUp garantiza conexión directa
        if proxy:
            env = {"NO_PROXY": "", "no_proxy": "", "ALL_PROXY": "", "all_proxy": "",
                   "HTTPS_PROXY": proxy, "https_proxy": proxy, "HTTP_PROXY": proxy, "http_proxy": proxy}
        mods = {"curl_cffi": None}
        if stdlib:
            mods["requests"] = None
        with mock.patch.dict(os.environ, env), mock.patch.dict(sys.modules, mods):
            F = macfetch.Fetch("t", self.tmp, budget_s=self.BUDGET, log=lambda *a: None)
            t0 = _t.monotonic()
            try:
                body, err = F.get(target_url, timeout=120), None
            except macfetch.FetchError as e:
                body, err = None, e
            return body, err, _t.monotonic() - t0, F

    def check_cut(self, body, err, elapsed, F, hits, aborted):
        self.assertLessEqual(elapsed, self.BUDGET + self.TOL, "%.3f s" % elapsed)
        self.assertIsNone(body)
        self.assertIsNotNone(err)
        self.assertNotEqual(F.requests_log[-1]["cls"], "OK")
        self.assertEqual(hits, 1)                                       # ni otro intento ni otro transporte
        self.assertTrue(aborted, "la conexión no se cortó")

    @NEEDS_REQUESTS
    def test_b3r2_1_https_via_proxy_slow_connect_cut_at_deadline(self):
        px = _Proxy("slow_connect")
        try:
            res = self.fetch_via("https://never-contacted.invalid/file", proxy=px.url)
            aborted = px.aborted.wait(3)
        finally:
            px.close()
        self.check_cut(*res, hits=px.hits, aborted=aborted)

    @NEEDS_REQUESTS
    def test_b3r2_1_https_via_proxy_tunnel_fast_response_ok(self):
        t = self.tls()
        srv, px = _Slow("fast", t), _Proxy("tunnel")
        try:
            body, err, elapsed, F = self.fetch_via(srv.url, proxy=px.url)
        finally:
            px.close()
            srv.close()
        self.assertIsNone(err)
        self.assertEqual(body, b"x" * 10)
        self.assertEqual(px.hits, 1)                                    # pasó de verdad por el túnel
        self.assertLess(elapsed, self.BUDGET)

    @NEEDS_REQUESTS
    def test_b3r2_1_https_via_proxy_tunnel_slow_body_cut_at_deadline(self):
        t = self.tls()
        srv, px = _Slow("drip", t), _Proxy("tunnel")
        try:
            res = self.fetch_via(srv.url, proxy=px.url)
            aborted = srv.aborted.wait(3)
        finally:
            px.close()
            srv.close()
        self.check_cut(*res, hits=srv.hits, aborted=aborted)
        self.assertEqual(px.hits, 1)

    @NEEDS_REQUESTS
    def test_b3r2_1_requests_slow_tls_handshake_cut_at_deadline(self):
        hs = _SlowTLSHandshake()
        try:
            res = self.fetch_via(hs.url)
            aborted = hs.aborted.wait(3)
        finally:
            hs.close()
        self.check_cut(*res, hits=hs.hits, aborted=aborted)

    def test_b3r2_1_stdlib_slow_tls_handshake_cut_at_deadline(self):
        hs = _SlowTLSHandshake()
        try:
            res = self.fetch_via(hs.url, stdlib=True)
            aborted = hs.aborted.wait(3)
        finally:
            hs.close()
        self.check_cut(*res, hits=hs.hits, aborted=aborted)

    @NEEDS_REQUESTS
    def test_proxy_connections_are_guarded_too(self):
        """Solo comprueba la configuración del gestor de proxy; la negociación real la cubren las pruebas b3r2_1."""
        import requests
        from g8common import g8http
        guard = g8http.SocketDeadline(60)
        self.addCleanup(guard.cancel)
        sess = macfetch._guarded_session(requests, guard)
        self.addCleanup(sess.close)
        a = sess.get_adapter("https://x")
        m = a.proxy_manager_for("http://proxy.invalid:3128")
        self.assertIs(m.pool_classes_by_scheme, a.poolmanager.pool_classes_by_scheme)
        self.assertTrue(issubclass(m.pool_classes_by_scheme["https"].ConnectionCls,
                                   __import__("urllib3").connection.HTTPSConnection))

    def test_no_guard_threads_left(self):
        import threading
        before = threading.active_count()
        self.call("fast")
        self.call("drip")
        time.sleep(0.2)
        self.assertLessEqual(threading.active_count(), before)


class WriteMetadata(unittest.TestCase):
    """B3-4 — max_date del registro = fecha máxima del CSV efectivamente escrito."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp)
        self.F = macfetch.Fetch("t", self.tmp, transport=lambda *a: None, log=lambda *a: None)
        self.p = os.path.join(self.tmp, "data", "X.csv")

    def check(self):
        written = S.parse(open(self.p, "rb").read())
        self.assertEqual(self.F.files["X.csv"]["max_date"], written.max_date)
        return self.F.files["X.csv"]

    def test_new_file(self):
        self.F.write(self.p, b"Date,Value\n2026-09-24,1.0\n", expect_header=["Date", "Value"])
        self.assertEqual(self.check()["max_date"], "20260924")

    def test_replacement_without_kept_rows(self):
        self.F.write(self.p, b"Date,Value\n2026-09-23,1.0\n", expect_header=["Date", "Value"])
        self.F.write(self.p, b"Date,Value\n2026-09-23,1.0\n2026-09-24,1.1\n", expect_header=["Date", "Value"])
        rec = self.check()
        self.assertEqual((rec["max_date"], rec["kept_from_previous"]), ("20260924", 0))

    def test_union_with_kept_history(self):
        self.F.write(self.p, b"Date,Value\n2026-09-20,1.0\n2026-09-25,1.2\n", expect_header=["Date", "Value"])
        self.F.write(self.p, b"Date,Value\n2026-09-24,1.1\n", expect_header=["Date", "Value"])
        rec = self.check()
        self.assertEqual((rec["max_date"], rec["kept_from_previous"]), ("20260925", 2))


class HeartbeatDetail(unittest.TestCase):
    """state/fetch_<clave>.json → latido del publicador → aviso de Actions con el motivo."""

    def test_detail_reaches_watch_alert(self):
        import ingest_watch as W
        root = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, root)
        os.makedirs(os.path.join(root, "sources"))
        open(os.path.join(root, "sources", "executors.csv"), "w").write(
            "executor_id,job,timezone,run_times_local,days,alert_after_min,active_from,recovery,provisional,justification\n"
            "mac-primary,nzchf-tona,Africa/Malabo,08:00,mon-sun,45,2026-09-20,x,Y,x\n")
        open(os.path.join(root, "sources", "credentials.csv"), "w").write("credential_id,expires_utc\n")
        p = os.path.join(root, "data", "_ingest", "latest", "mac-primary__nzchf-tona.json")
        os.makedirs(os.path.dirname(p))
        json.dump({"run_id": "r", "started_utc": "2026-09-24T07:00:05Z", "families": {}, "credentials": {"days_left": 90},
                   "fetch_status": {"chf": "1"},
                   "fetch_detail": {"chf": {"errors": ["saron: download failed: FAIL_TRANSIENT HTTP 503"]}}}, open(p, "w"))
        a, info = {}, []
        W.check_executors(root, datetime(2026, 9, 24, 8, 0).astimezone(), a, info)
        self.assertIn("saron: download failed", a["exec:mac-primary:nzchf-tona:fetch:chf"])

    def test_publisher_attaches_fetch_detail(self):
        import test_mac_push as TP
        t = TP.MacPushTests("test_no_news_heartbeat_only")
        t.setUp()
        try:
            os.makedirs(os.path.join(t.tmp, "state"), exist_ok=True)
            json.dump({"rc": 1, "errors": ["descarga B2: FAIL_ACCESS 403"], "files": {}, "requests": []},
                      open(os.path.join(t.tmp, "state", "fetch_nzd.json"), "w"))
            rc, rec, _ = t.run_push(t.fresh_fetch()[:1] + ["nzd=1,chf=0,tona=0"] + t.fresh_fetch()[2:])
            self.assertIn("FAIL_ACCESS", rec["fetch_detail"]["nzd"]["errors"][0])
            self.assertIn("FAIL_ACCESS", json.dumps(rec["alerts"]) + json.dumps(rec))
        finally:
            t.tearDown()


if __name__ == "__main__":
    unittest.main()

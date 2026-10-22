"""B6: un IIB neozelandés vencido no se reescribe. El CSV que ya está en data/ queda cerrado.

Sin red. El libro B2 se fabrica en memoria. La fecha de la pasada es el calendario de Auckland.
"""
import csv
import hashlib
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from datetime import date, datetime, timezone

from openpyxl import Workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import fetch_nzd_b2 as B  # noqa: E402
import real_yields_g8 as RY  # noqa: E402
from g8common.freshness import evaluate  # noqa: E402

# Historia ya publicada. Este lote no puede mover estos bytes.
IIB_2025_SHA = "ee316afdc5ea236f011858959a119fbcde40ae55a2ecb154883199449e28f484"
S01B_SHA = "0c43d6dd54677974bf4ffde880de049396f4b0d9cdd1c6e39964d796ef132ee5"
ACM_NZD_SHA = "47f9049859fb7aa7e86c543f1307d9e189b06038ff71d98595c9ef8dd31b2640"

CORE = (
    ("INM.DB01.NZZV", "NZD_BILL_30D.csv"),
    ("INM.DB02.NZZV", "NZD_BILL_60D.csv"),
    ("INM.DB03.NZZV", "NZD_BILL_90D.csv"),
    ("INM.DG101.NZZCF", "NZD_BOND_1Y.csv"),
    ("INM.DG102.NZZCF", "NZD_BOND_2Y.csv"),
    ("INM.DG105.NZZCF", "NZD_BOND_5Y.csv"),
    ("INM.DG110.NZZCF", "NZD_BOND_10Y.csv"),
    ("INM.DN.NZK", "NZD_CASH_ON.csv"),
)


def _sha(path):
    with open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def _rules(name):
    with open(os.path.join(ROOT, "sources", name), newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _workbook(iibs):
    """iibs: [(year, [(YYYY-MM-DD, float), ...]), ...]. El resto de la B2 lleva un cierre."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.cell(1, 1, "Series")
    cols = []
    for sid, fname in CORE:
        cols.append((sid, "Wholesale rates", fname[:-4], [("2026-10-15", 4.25)]))
    for year, rows in iibs:
        cols.append(("INM.DI%d.NZZC" % year, "Inflation indexed bonds", "20 September %d" % year, rows))
    for j, (sid, label, sub, _rows) in enumerate(cols, start=2):
        ws.cell(1, j, label)
        ws.cell(2, j, sub)
        ws.cell(5, j, sid)
    dates = sorted({d for *_, rows in cols for d, _v in rows})
    look = []
    for _sid, _label, _sub, rows in cols:
        look.append({d: v for d, v in rows})
    for i, d in enumerate(dates, start=6):
        ws.cell(i, 1, datetime.strptime(d, "%Y-%m-%d"))
        for j, series in enumerate(look, start=2):
            if d in series:
                ws.cell(i, j, float(series[d]))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class MaturityRule(unittest.TestCase):
    def test_closed_the_day_after_20_september(self):
        self.assertFalse(B.iib_is_closed("NZD_IIB_2025.csv", date(2025, 9, 20)))
        self.assertTrue(B.iib_is_closed("NZD_IIB_2025.csv", date(2025, 9, 21)))
        self.assertTrue(B.iib_is_closed("NZD_IIB_2025.csv", date(2026, 10, 17)))
        self.assertFalse(B.iib_is_closed("NZD_IIB_2030.csv", date(2026, 10, 17)))
        self.assertFalse(B.iib_is_closed("NZD_IIB_2035.csv", date(2026, 10, 17)))
        self.assertFalse(B.iib_is_closed("NZD_BOND_10Y.csv", date(2036, 1, 1)))
        self.assertEqual(B.iib_maturity(2025), date(2025, 9, 20))

    def test_auckland_calendar_not_utc(self):
        # 17-oct-2026 10:00Z sigue siendo 17 en Auckland; 12:00Z ya es 18 (NZDT, UTC+13).
        self.assertEqual(B.auckland_today(datetime(2026, 10, 17, 10, 0, tzinfo=timezone.utc)), date(2026, 10, 17))
        self.assertEqual(B.auckland_today(datetime(2026, 10, 17, 12, 0, tzinfo=timezone.utc)), date(2026, 10, 18))
        self.assertEqual(B.auckland_today(datetime(2025, 9, 20, 23, 30)), date(2025, 9, 20))


class DailyPassLeavesClosedHistory(unittest.TestCase):
    def setUp(self):
        self.map_before = dict(B.SERIES_MAP)
        B._F = None
        self.roots = []

    def tearDown(self):
        B.SERIES_MAP.clear()
        B.SERIES_MAP.update(self.map_before)
        B._F = None
        for root in self.roots:
            shutil.rmtree(root, ignore_errors=True)

    def _root(self):
        root = tempfile.mkdtemp()
        os.makedirs(os.path.join(root, "data"))
        self.roots.append(root)
        return root

    def _run(self, root, payload, today):
        B._F = None
        B.SERIES_MAP.clear()
        B.SERIES_MAP.update(self.map_before)

        def _download(url):
            return payload

        old_dl, old_today = B.download_xlsx, B.auckland_today
        old_argv, old_cwd = sys.argv, os.getcwd()
        B.download_xlsx = _download
        B.auckland_today = lambda now=None: today
        sys.argv = ["fetch_nzd_b2.py", "--no-hist"]
        os.chdir(root)
        out, err = io.StringIO(), io.StringIO()
        try:
            with redirect_stdout(out), redirect_stderr(err):
                rc = B.main()
        finally:
            B.download_xlsx, B.auckland_today = old_dl, old_today
            sys.argv, os.chdir(old_cwd)
            B._F = None
        return rc, out.getvalue(), err.getvalue()

    def test_matured_csv_is_not_rewritten_and_live_iib_is(self):
        root = self._root()
        closed = os.path.join(root, "data", "NZD_IIB_2025.csv")
        seed = b"Date,Value\n2025-10-20,0.88\n"
        with open(closed, "wb") as fh:
            fh.write(seed)
        stamp = 1_700_000_000
        os.utime(closed, (stamp, stamp))
        payload = _workbook([
            (2024, [("2024-09-01", 1.5)]),          # vencido y sin CSV: no se crea
            (2025, [("2026-10-15", 9.99)]),         # la columna sigue en el xlsx, con otro número
            (2035, [("2026-10-15", 2.69)]),
        ])
        rc, log, err = self._run(root, payload, date(2026, 10, 17))
        self.assertEqual(rc, 0, err + log)
        with open(closed, "rb") as fh:
            self.assertEqual(fh.read(), seed)
        self.assertEqual(os.stat(closed).st_mtime, stamp)
        self.assertFalse(os.path.exists(os.path.join(root, "data", "NZD_IIB_2024.csv")))
        with open(os.path.join(root, "data", "NZD_IIB_2035.csv"), encoding="utf-8") as fh:
            live = fh.read()
        self.assertIn("2026-10-15,2.69", live)
        with open(os.path.join(root, "data", "NZD_BOND_10Y.csv"), encoding="utf-8") as fh:
            bond = fh.read()
        self.assertIn("2026-10-15,4.25", bond)
        self.assertIn("NZD_IIB_2025.csv: historia cerrada (vencimiento 2025-09-20)", log)
        self.assertIn("NZD_IIB_2024.csv: historia cerrada (vencimiento 2024-09-20)", log)
        with open(os.path.join(root, "state", "fetch_nzd.json"), encoding="utf-8") as fh:
            state = json.load(fh)
        self.assertEqual(state["rc"], 0)
        self.assertNotIn("NZD_IIB_2025.csv", state["files"])
        self.assertEqual(state["files"]["NZD_IIB_2035.csv"]["status"], "WRITTEN")
        self.assertEqual(state["files"]["NZD_BOND_10Y.csv"]["status"], "WRITTEN")

    def test_maturity_day_still_writes_the_last_close(self):
        root = self._root()
        payload = _workbook([(2025, [("2025-09-20", 1.11)]), (2035, [("2025-09-20", 2.22)])])
        rc, log, err = self._run(root, payload, date(2025, 9, 20))
        self.assertEqual(rc, 0, err + log)
        with open(os.path.join(root, "data", "NZD_IIB_2025.csv"), encoding="utf-8") as fh:
            text = fh.read()
        self.assertIn("2025-09-20,1.11", text)
        self.assertNotIn("historia cerrada", log)


class RealYieldAndFreshness(unittest.TestCase):
    def test_closed_file_remains_in_the_interpolation(self):
        """build_nzd sigue leyendo todo NZD_IIB_*.csv. Un 2025 que ya no imprime no tapa el día vivo."""
        folder = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, folder)

        def put(name, rows):
            with open(os.path.join(folder, name), "w", encoding="utf-8") as fh:
                fh.write("Date,Value\n" + "".join("%s,%s\n" % row for row in rows))

        put("NZD_IIB_2025.csv", [("2024-09-20", "1.00")])
        put("NZD_IIB_2035.csv", [("2024-09-20", "3.00"), ("2026-10-15", "2.50")])
        put("NZD_BOND_10Y.csv", [("2024-09-20", "4.00"), ("2026-10-15", "4.20")])
        with open(os.path.join(ROOT, "scripts", "real_yields_g8.py"), encoding="utf-8") as fh:
            src = fh.read()
        self.assertIn('glob.glob(os.path.join(DATA_DIR, "NZD_IIB_*.csv"))', src)
        self.assertNotIn("iib_is_closed", src)
        old = RY.DATA_DIR
        RY.DATA_DIR = folder
        try:
            _nom, real, _xchk, quality, desc = RY.build_nzd()
        finally:
            RY.DATA_DIR = old
        self.assertEqual(quality, "PROXY_THIN_MARKET")
        self.assertIn("2025", desc)
        self.assertIn("2035", desc)
        self.assertEqual(str(real.index[-1].date()), "2026-10-15")
        self.assertAlmostEqual(float(real.loc["2026-10-15"]), 2.50, places=6)
        self.assertIn(date(2024, 9, 20), {ts.date() for ts in real.index})

    def test_freshness_keeps_2025_ended_and_out_of_the_live_rule(self):
        rules = {r["rule_id"]: r for r in _rules("freshness_rules.csv")}
        params = {r["rule_id"]: r for r in _rules("freshness_params.csv")}
        ended = rules["NZ_IIB_2025"]
        self.assertEqual(ended["files"], "NZD_IIB_2025.csv")
        self.assertEqual(ended["row_class"], "input_ended")
        self.assertEqual(ended["frequency"], "ended")
        self.assertEqual(ended["status"], "SIN_FRESCURA")
        self.assertNotIn("NZD_IIB_2025.csv", rules["NZ_B2"]["files"].split(";"))
        for live in ("NZD_IIB_2030.csv", "NZD_IIB_2035.csv", "NZD_IIB_2040.csv", "NZD_IIB_2050.csv"):
            self.assertIn(live, rules["NZ_B2"]["files"])
        ry = rules["RY_NZD"]
        self.assertNotIn("NZD_IIB_2025.csv", ry["files"].split(";"))
        self.assertEqual(params["NZ_IIB_2025"]["model"], "not_monitored")
        res = evaluate(
            {"file": "NZD_IIB_2025.csv", "calendar": "NZ", "timezone": "Pacific/Auckland"},
            ended,
            params["NZ_IIB_2025"],
            {},
            datetime(2026, 10, 17, tzinfo=timezone.utc),
            [date(2025, 10, 20)],
        )
        self.assertEqual(res["state"], "NOT_MONITORED")

    def test_published_history_bytes_stay(self):
        self.assertEqual(_sha(os.path.join(ROOT, "data", "NZD_IIB_2025.csv")), IIB_2025_SHA)
        self.assertEqual(_sha(os.path.join(ROOT, "scripts", "s01b.py")), S01B_SHA)
        with open(os.path.join(ROOT, "data", "ACM_G8_NZD.csv"), encoding="utf-8") as fh:
            acm_lines = fh.read().splitlines()
        self.assertEqual(acm_lines[0], "DATE,Y10_FIT,RNY10,TP10,QUALITY")
        self.assertGreater(len(acm_lines), 1000)
        with open(os.path.join(ROOT, "data", "NZD_IIB_2025.csv"), encoding="utf-8") as fh:
            last = fh.read().splitlines()[-1]
        self.assertEqual(last, "2025-10-20,0.88")


if __name__ == "__main__":
    unittest.main()

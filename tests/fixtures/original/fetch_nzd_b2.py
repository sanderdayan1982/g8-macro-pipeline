"""
fetch_nzd_b2.py — RBNZ Table B2 (Wholesale interest rates) daily close  — v1.4 (2026-09-13)

Sources NZD short-end and government bond yields from a single RBNZ XLSX:
  - BKBM 30/60/90D bank bill yields (NZD analogue of BBSW/SOFR-bills)
  - NZ Government Bonds 1Y/2Y/5Y/10Y constant maturity closing yields
  - Overnight interbank cash rate (INM.DN.NZK) → NZD_CASH_ON.csv  (§02 floor NZD = cash − OCR)

Endpoint:
  https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily-close.xlsx

Outputs (one CSV per series, Date,Value with YYYY-MM-DD dates) into ./data/.

Network (v1.2+): the RBNZ WAF rejects python-requests' TLS fingerprint even from
residential IPs (403). We download with curl_cffi impersonating Safari when
available (pip install --user curl_cffi) and fall back to requests otherwise.
Runs locally on the operator's Mac via launchd (Trading_Sander/g8-nzd);
push_nzd_to_github.py uploads the CSVs to data/ through the GitHub REST API.
NOT run by GitHub Actions (datacenter IPs are blocked by the RBNZ WAF).

Column matching is done by Series Id (row 5 of the Data sheet), NOT by column position.
`--list` prints every column's header rows (1–5) to pin new Series Ids.

v1.4 (NZD real — same formula as the rest of the G8):
  - Inflation-indexed bonds (IIB): every column whose header mentions "inflation"/"indexed"
    and a maturity year is written as NZD_IIB_<year>.csv (real yields → real_yields_g8.py
    builds RY_G8_NZD.csv: REAL10 = IIB interpolated to 10Y, BE10 = NOM10 − REAL10).
  - History splice: the static 1985-2017 daily-close workbook (HIST_URL) is downloaded ONCE
    (cached in data/_cache/) and prepended to bills/bonds/cash so the CSVs carry 40 years —
    enough monthly obs for acm_g8.py to estimate a real ACM (K=3) for NZD instead of the
    AUD-anchored proxy. Series Ids are matched first; label fallback (HIST_LABEL_FALLBACK)
    covers a schema change in the old file. `--list-hist` prints the old file's columns.
    `--no-hist` skips the splice (current file only, v1.3 behaviour).
"""

import os
import sys
import time
from io import BytesIO
from openpyxl import load_workbook

try:
    from curl_cffi import requests as crequests      # TLS impersonation
    _IMPERSONATE = True
except Exception:                                    # pragma: no cover
    crequests = None
    _IMPERSONATE = False
import requests

URL = (
    "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/"
    "statistics/series/b/b2/hb2-daily-close.xlsx"
)
HIST_URL = (
    "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/"
    "statistics/series/b/b2/hb2-daily-close-1985-2017.xlsx"
)
REFERER = "https://www.rbnz.govt.nz/statistics/series/exchange-and-interest-rates/wholesale-interest-rates"

HEADERS = {
    "User-Agent": ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                   "AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.0 Safari/605.1.15"),
    "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": REFERER,
}

OUTPUT_DIR = "data"
TIMEOUT_SECONDS = 90
RETRY_DELAYS = [0, 5, 15]

SERIES_MAP = {
    "INM.DB01.NZZV":   "NZD_BILL_30D.csv",
    "INM.DB02.NZZV":   "NZD_BILL_60D.csv",
    "INM.DB03.NZZV":   "NZD_BILL_90D.csv",
    "INM.DG101.NZZCF": "NZD_BOND_1Y.csv",
    "INM.DG102.NZZCF": "NZD_BOND_2Y.csv",
    "INM.DG105.NZZCF": "NZD_BOND_5Y.csv",
    "INM.DG110.NZZCF": "NZD_BOND_10Y.csv",
    # v1.3: pinned from `--list` on 2026-09-13 (col 5: Cash rate | Overnight interbank cash rate | %pa)
    "INM.DN.NZK":      "NZD_CASH_ON.csv",
}
# Label discovery disabled: "overnight" also matches the Overnight Deposit Rate (INM.DD1.N),
# which is the RBNZ floor, not the traded cash rate. Pin by Series Id only.
OPTIONAL_BY_LABEL = {}

# v1.4: inflation-indexed bonds — discovered by header text (joined rows 1-5 of the column):
# must mention inflation/indexed AND a 4-digit maturity year (e.g. "Inflation indexed bonds |
# 20 September 2035 | %"). One CSV per maturity: NZD_IIB_2035.csv. Loud WARN if none found.
IIB_PATTERN = r"(inflation|indexed|index-linked|IIB)"
IIB_YEAR = r"(20[2-6]\d)"

# v1.4: label fallback for the 1985-2017 workbook if its Series Ids differ from the current one.
# Keys are regexes tested against the joined header text; values are output files.
HIST_LABEL_FALLBACK = {
    r"bank\s*bill.*\b30\b|\b30\s*day": "NZD_BILL_30D.csv",
    r"bank\s*bill.*\b60\b|\b60\s*day": "NZD_BILL_60D.csv",
    r"bank\s*bill.*\b90\b|\b90\s*day": "NZD_BILL_90D.csv",
    r"government\s*bond.*\b1\s*year|\b1\s*year.*bond": "NZD_BOND_1Y.csv",
    r"government\s*bond.*\b2\s*year|\b2\s*year.*bond": "NZD_BOND_2Y.csv",
    r"government\s*bond.*\b5\s*year|\b5\s*year.*bond": "NZD_BOND_5Y.csv",
    r"government\s*bond.*\b10\s*year|\b10\s*year.*bond": "NZD_BOND_10Y.csv",
    r"overnight\s*interbank|interbank\s*cash": "NZD_CASH_ON.csv",
}
CACHE_DIR = os.path.join("data", "_cache")

SHEET_NAME = "Data"
SERIES_ID_ROW = 5
DATA_START_ROW = 6
DATE_COL = 1


def _get(url):
    """One HTTP GET: curl_cffi (Safari TLS) first, plain requests as fallback."""
    if _IMPERSONATE:
        for imp in ("safari17_0", "chrome124", "chrome"):
            try:
                r = crequests.get(url, headers=HEADERS, timeout=TIMEOUT_SECONDS, impersonate=imp)
                print(f"[fetch_nzd_b2] curl_cffi impersonate={imp} → HTTP {r.status_code}")
                if r.status_code == 200:
                    return r.content
            except Exception as e:
                print(f"[fetch_nzd_b2] curl_cffi {imp} error: {e}")
    r = requests.get(url, timeout=TIMEOUT_SECONDS, headers=HEADERS)
    print(f"[fetch_nzd_b2] requests → HTTP {r.status_code}")
    r.raise_for_status()
    return r.content


def download_xlsx(url: str) -> bytes:
    last_exc = None
    for attempt, delay in enumerate(RETRY_DELAYS, start=1):
        if delay:
            print(f"[fetch_nzd_b2] retry in {delay}s...")
            time.sleep(delay)
        try:
            print(f"[fetch_nzd_b2] attempt {attempt}/{len(RETRY_DELAYS)} GET {url}")
            content = _get(url)
            print(f"[fetch_nzd_b2] downloaded {len(content):,} bytes")
            return content
        except Exception as e:
            last_exc = e
            print(f"[fetch_nzd_b2] error: {e}")
            if "403" in str(e):
                break                       # WAF verdict: retrying won't help
    raise RuntimeError(f"[fetch_nzd_b2] download failed: {last_exc}")


def locate_series_columns(ws) -> dict:
    found = {}
    for row in ws.iter_rows(min_row=SERIES_ID_ROW, max_row=SERIES_ID_ROW, values_only=False):
        for cell in row:
            if isinstance(cell.value, str) and cell.value in SERIES_MAP:
                found[cell.value] = cell.column
    missing = set(SERIES_MAP.keys()) - set(found.keys())
    if missing:
        raise RuntimeError(f"[fetch_nzd_b2] missing Series Ids in XLSX (RBNZ schema changed?): {sorted(missing)}")
    import re as _re
    headers = {}
    for row in ws.iter_rows(min_row=1, max_row=SERIES_ID_ROW, values_only=False):
        for cell in row:
            if isinstance(cell.value, str):
                headers.setdefault(cell.column, []).append(cell.value)
    for pattern, fname in OPTIONAL_BY_LABEL.items():
        hit = None
        for col, labels in headers.items():
            if any(_re.search(pattern, lb, _re.I) for lb in labels) and col not in found.values():
                hit = col
                break
        if hit is None:
            print(f"[fetch_nzd_b2] WARN optional series /{pattern}/ not found — {fname} not written")
        else:
            sid = "LABEL:" + fname
            SERIES_MAP[sid] = fname
            found[sid] = hit
            print(f"[fetch_nzd_b2] optional /{pattern}/ → column {hit} ({' | '.join(headers[hit])[:80]}) → {fname}")
    # v1.4: inflation-indexed bonds by header text (one file per maturity year)
    n_iib = 0
    for col, labels in sorted(headers.items()):
        if col in found.values():
            continue
        text = " | ".join(labels)
        if not _re.search(IIB_PATTERN, text, _re.I):
            continue
        yr = _re.findall(IIB_YEAR, text)
        if not yr:
            continue
        fname = f"NZD_IIB_{yr[-1]}.csv"
        sid = "IIB:" + fname
        if sid in SERIES_MAP:
            continue                                     # first column wins (yield before price)
        SERIES_MAP[sid] = fname
        found[sid] = col
        n_iib += 1
        print(f"[fetch_nzd_b2] IIB column {col} ({text[:80]}) → {fname}")
    if n_iib == 0:
        print("[fetch_nzd_b2] WARN no inflation-indexed bond columns found — run --list and pin them; "
              "RY_G8_NZD.csv will stay on the manual BE constant")
    return found


def list_columns(ws) -> None:
    cols = {}
    for row in ws.iter_rows(min_row=1, max_row=SERIES_ID_ROW, values_only=False):
        for cell in row:
            if cell.value is not None:
                cols.setdefault(cell.column, []).append(str(cell.value))
    for col in sorted(cols):
        print(f"col {col:3d}: " + " | ".join(cols[col]))


def extract_series(ws, col_map: dict) -> dict:
    out = {sid: [] for sid in col_map}
    targets = [(sid, col - 1) for sid, col in col_map.items()]
    date_idx = DATE_COL - 1
    for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        if not row or row[date_idx] is None:
            continue
        try:
            date_str = row[date_idx].strftime("%Y-%m-%d")
        except AttributeError:
            continue
        for sid, idx in targets:
            if idx >= len(row):
                continue
            v = row[idx]
            if v is None or v == "":
                continue
            try:
                out[sid].append((date_str, float(v)))
            except (TypeError, ValueError):
                continue
    return out


def write_csv(path: str, rows) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write("Date,Value\n")
        for date_str, value in rows:
            f.write(f"{date_str},{round(value, 6)}\n")


def hist_workbook():
    """1985-2017 daily-close workbook: downloaded once, cached on disk (static file)."""
    os.makedirs(CACHE_DIR, exist_ok=True)
    cp = os.path.join(CACHE_DIR, os.path.basename(HIST_URL))
    if os.path.exists(cp) and os.path.getsize(cp) > 50_000:
        print(f"[fetch_nzd_b2] hist: using cached {cp}")
        content = open(cp, "rb").read()
    else:
        content = download_xlsx(HIST_URL)
        with open(cp, "wb") as f:
            f.write(content)
        print(f"[fetch_nzd_b2] hist: cached → {cp}")
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
    return ws


def hist_series(ws) -> dict:
    """Extract the same target series from the old workbook → {filename: [(date, v)]}.
    Series Id first, header-text fallback second; never raises (WARN + partial)."""
    import re as _re
    headers = {}
    for row in ws.iter_rows(min_row=1, max_row=SERIES_ID_ROW, values_only=False):
        for cell in row:
            if isinstance(cell.value, str):
                headers.setdefault(cell.column, []).append(cell.value)
    ids = {}
    for row in ws.iter_rows(min_row=SERIES_ID_ROW, max_row=SERIES_ID_ROW, values_only=False):
        for cell in row:
            if isinstance(cell.value, str):
                ids[cell.value.strip()] = cell.column
    col_by_file = {}
    for sid, fname in SERIES_MAP.items():
        if sid in ids and not sid.startswith(("LABEL:", "IIB:")):
            col_by_file[fname] = ids[sid]
    for pattern, fname in HIST_LABEL_FALLBACK.items():
        if fname in col_by_file:
            continue
        for col, labels in sorted(headers.items()):
            if _re.search(pattern, " | ".join(labels), _re.I) and col not in col_by_file.values():
                col_by_file[fname] = col
                print(f"[fetch_nzd_b2] hist: label fallback /{pattern}/ → column {col} → {fname}")
                break
    missing = [f for f in HIST_LABEL_FALLBACK.values() if f not in col_by_file]
    if missing:
        print(f"[fetch_nzd_b2] WARN hist: not found in 1985-2017 workbook: {missing} (run --list-hist)")
    raw = extract_series(ws, {f: c for f, c in col_by_file.items()})
    return raw


def splice(current_rows, hist_rows):
    """Prepend history strictly before the first current date; both sorted by date."""
    if not hist_rows:
        return current_rows
    first = current_rows[0][0] if current_rows else "9999-12-31"
    old = [r for r in sorted(hist_rows) if r[0] < first]
    return old + current_rows


def main() -> int:
    if "--list-hist" in sys.argv:
        try:
            list_columns(hist_workbook())
            return 0
        except Exception as e:
            print(f"[fetch_nzd_b2] FATAL hist: {e}", file=sys.stderr)
            return 1
    try:
        content = download_xlsx(URL)
    except Exception as e:
        print(f"[fetch_nzd_b2] FATAL download: {e}", file=sys.stderr)
        return 1
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
        if SHEET_NAME not in wb.sheetnames:
            raise RuntimeError(f"sheet '{SHEET_NAME}' not found. Sheets: {wb.sheetnames}")
        ws = wb[SHEET_NAME]
        if "--list" in sys.argv:
            list_columns(ws)
            return 0
        col_map = locate_series_columns(ws)
        print(f"[fetch_nzd_b2] located all {len(col_map)} target series")
        data = extract_series(ws, col_map)
    except Exception as e:
        print(f"[fetch_nzd_b2] FATAL parse: {e}", file=sys.stderr)
        return 2
    # v1.4: splice 1985-2017 history (cached workbook) — never fatal
    hist = {}
    if "--no-hist" not in sys.argv:
        try:
            hist = hist_series(hist_workbook())
            print(f"[fetch_nzd_b2] hist: {len(hist)} series, "
                  + ", ".join(f"{k} {len(v):,}" for k, v in hist.items()))
        except Exception as e:
            print(f"[fetch_nzd_b2] WARN hist splice skipped: {e}")
    total = 0
    for sid, rows in data.items():
        filename = SERIES_MAP[sid]
        rows = splice(rows, hist.get(filename))
        write_csv(os.path.join(OUTPUT_DIR, filename), rows)
        ld, lv = rows[-1] if rows else ("NONE", float("nan"))
        print(f"[fetch_nzd_b2] {filename}: {len(rows):,} rows, latest {ld} = {lv}")
        total += len(rows)
    print(f"[fetch_nzd_b2] OK — {len(data)} files, {total:,} total rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
